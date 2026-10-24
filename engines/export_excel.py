"""Lazy Excel export (split from export_engine bloat). openpyxl.

Enterprise skin: brand headers, severity fills, currency formats, frozen
panes, autofilters, sized columns.
"""

import io

_BRAND_FILL = "0071E3"
_HEADER_FONT = {"bold": True, "color": "FFFFFF", "size": 10, "name": "Google Sans Flex"}
_SEV_FILLS = {
    "CRITICAL": "F8D7DA",
    "HIGH": "FDE7D3",
    "MEDIUM": "FBF0CE",
    "LOW": "D9F2E3",
    "INFO": "DCE9FB",
}
_SEV_FONT_COLORS = {
    "CRITICAL": "B00020",
    "HIGH": "C44900",
    "MEDIUM": "8A6D00",
    "LOW": "0F6B33",
    "INFO": "1D4ED8",
}


def _style_header(ws, headers, widths=None):
    from openpyxl.styles import Alignment, Font, PatternFill

    for idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=idx)
        cell.fill = PatternFill("solid", fgColor=_BRAND_FILL)
        cell.font = Font(**_HEADER_FONT)
        cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    if headers:
        ws.auto_filter.ref = f"A1:{chr(64 + min(len(headers), 26))}1" if len(headers) <= 26 else None
    if widths:
        from openpyxl.utils import get_column_letter

        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def _style_severity_cell(ws, row, col):
    from openpyxl.styles import Font, PatternFill

    cell = ws.cell(row=row, column=col)
    sev = str(cell.value or "").upper()
    if sev in _SEV_FILLS:
        cell.fill = PatternFill("solid", fgColor=_SEV_FILLS[sev])
        cell.font = Font(bold=True, color=_SEV_FONT_COLORS[sev], name="Google Sans Flex")


def _currency_cols(ws, cols, first_row=2):
    for col in cols:
        for row in range(first_row, ws.max_row + 1):
            ws.cell(row=row, column=col).number_format = '"$"#,##0'


def build_excel_workbook(line_items: list, rooms: list = None, comps: dict = None) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    wb.properties.title = "Repair Estimate - Line Items"
    wb.properties.creator = "RealEstate Repair Cost Estimator"

    ws = wb.active
    ws.title = "Line items"
    ws.sheet_properties.tabColor = _BRAND_FILL
    ws.append(["System", "Severity", "Finding", "Low", "High", "Provenance"])
    for it in line_items or []:
        ws.append(
            [
                it.get("system"),
                it.get("severity"),
                str(it.get("finding", ""))[:200],
                it.get("contractor_low", it.get("total_low")),
                it.get("contractor_high", it.get("total_high")),
                it.get("provenance"),
            ]
        )
    _style_header(
        ws, ["System", "Severity", "Finding", "Low", "High", "Provenance"], [14, 12, 78, 12, 12, 16]
    )
    _currency_cols(ws, (4, 5))
    for row in range(2, ws.max_row + 1):
        _style_severity_cell(ws, row, 2)

    if rooms:
        w2 = wb.create_sheet("Rooms")
        w2.sheet_properties.tabColor = _BRAND_FILL
        headers2 = ["Room", "Condition", "Priority", "Low", "Mid", "High", "Top issue"]
        w2.append(headers2)
        for r in rooms:
            w2.append(
                [
                    r.get("room"),
                    r.get("condition"),
                    r.get("priority"),
                    r.get("low"),
                    r.get("mid"),
                    r.get("high"),
                    r.get("top_issue"),
                ]
            )
        _style_header(w2, headers2, [16, 14, 12, 12, 12, 12, 60])
        _currency_cols(w2, (4, 5, 6))
        for row in range(2, w2.max_row + 1):
            _style_severity_cell(w2, row, 3)

    if comps and comps.get("comps_weighted"):
        w3 = wb.create_sheet("Comps")
        w3.sheet_properties.tabColor = _BRAND_FILL
        headers3 = ["Price", "Sqft", "$/sqft", "Dist mi", "Recency d", "Weight"]
        w3.append(headers3)
        for c in comps["comps_weighted"]:
            w3.append(
                [
                    c.get("price"),
                    c.get("sqft"),
                    c.get("ppsf"),
                    c.get("distance_mi"),
                    c.get("recency_days"),
                    c.get("weight_norm"),
                ]
            )
        _style_header(w3, headers3, [14, 10, 12, 10, 12, 10])
        _currency_cols(w3, (1, 3))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
